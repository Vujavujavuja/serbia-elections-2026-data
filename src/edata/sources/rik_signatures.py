"""RIK: statistics of certified supporter statements (signatures) per candidate list.

Every list must submit 10,000 certified voter statements (fewer for national minorities).
RIK publishes, per list, a spreadsheet of how the statements were certified on
https://www.rik.parlament.gov.rs/tekst/101175/broj-overenih-izjava-biraca.php
That page holds files for several elections; all are stored. Spreadsheets are kept as
files and parsed cell by cell, since their layouts differ between elections.
"""
from __future__ import annotations

import io
import re
from urllib.parse import quote, unquote, urljoin

import openpyxl
import xlrd
from lxml import html as lxml_html

from ..http import Fetcher
from ..store import RawStore
from ..tables import write_table
from ._common import provenance

NAME = "rik_signatures"
DESCRIPTION = "RIK spreadsheets of certified supporter signatures per candidate list"
PAGE = "https://www.rik.parlament.gov.rs/tekst/101175/broj-overenih-izjava-biraca.php"


def fetch(store: RawStore) -> None:
    fetcher = Fetcher(store, NAME, min_interval=0.5, timeout=300)
    _, content = fetcher.fetch(PAGE, ext=".html", tags={"kind": "page"})
    have = {r["tags"].get("file_url") for r in store.manifest(NAME) if r["status"] == 200}
    doc = lxml_html.fromstring(content.decode("utf-8", "replace"))
    for position, link in enumerate(doc.xpath("//a[@href]")):
        href = link.get("href").strip()
        if "/extfile/" not in href:
            continue
        url = urljoin(PAGE, quote(unquote(href), safe="/:%"))
        if url in have:
            continue
        have.add(url)
        fetcher.fetch(url, ext="." + href.rsplit(".", 1)[-1].lower(),
                      tags={"kind": "file", "file_url": url, "file_id": href.split("/")[3],
                            "file_name": unquote(href.rsplit("/", 1)[-1]), "position_on_page": position})
    fetcher.close()


def _sheets(name: str, content: bytes):
    if name.lower().endswith(".xlsx"):
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        for sheet in book.worksheets:
            yield sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]
    elif name.lower().endswith(".xls"):
        book = xlrd.open_workbook(file_contents=content)
        for sheet in book.sheets():
            yield sheet.name, [sheet.row_values(i) for i in range(sheet.nrows)]


def parse(store: RawStore) -> dict[str, int]:
    files, cells = [], []
    for url, r in store.latest(NAME, key=lambda r: r["tags"].get("file_url")).items():
        t = r["tags"]
        if t.get("kind") != "file":
            continue
        files.append({"file_id": t["file_id"], "file_name": t["file_name"],
                      "position_on_page": t["position_on_page"], "size": r["size"],
                      "local_path": r["path"], **provenance(r)})
        try:
            for sheet, rows in _sheets(t["file_name"], store.read(r)):
                for i, row in enumerate(rows):
                    for j, value in enumerate(row):
                        if value not in (None, ""):
                            cells.append({"file_id": t["file_id"], "file_name": t["file_name"], "sheet": sheet,
                                          "row": i, "col": j, "value": str(value), "raw_sha256": r["sha256"]})
        except Exception as exc:  # keep going; the file itself is in the raw store
            files[-1]["parse_error"] = repr(exc)
    write_table(NAME, "files", files)
    write_table(NAME, "cells", cells)
    return {"files": len(files), "cells": len(cells)}
